"""Admin reports, transactions, promo codes, and data downloads."""
import os, csv, json, random, string
from io import BytesIO, StringIO
from datetime import datetime
from zoneinfo import ZoneInfo
import pandas as pd
import psycopg2
from openpyxl import Workbook
from datetime import time as dttime
from flask import Blueprint, render_template, redirect, url_for, flash, request, g, send_file
from flask_login import login_required, current_user
from app.extensions import db
from app.models.product import Product, ProductSize
from app.models.order import Order, OrderItem
from app.models.payment import Transaction, PromoCode
from app.utils.helpers import today, mydtfmt, end_date_correction
from app.utils.filters import datetimeformat, to_cairo_time

cairo_tz = ZoneInfo("Africa/Cairo")
admin_reports_bp = Blueprint('admin_reports', __name__)


@admin_reports_bp.route("/admin/reports", methods=["GET"], endpoint='admin_reports')
@login_required
def admin_reports():
    if current_user.role not in ['admin', 'superadmin']: return redirect(url_for("home"))
    page = request.args.get('page', 1, type=int)
    start_date = request.args.get("start_date", None); end_date = request.args.get("end_date", None)
    status = request.args.get("status", ""); payment_method = request.args.get("payment_method", "")
    if start_date: start_date = mydtfmt(start_date).replace(tzinfo=cairo_tz).astimezone(ZoneInfo("UTC"))
    if end_date: end_date = mydtfmt(end_date).replace(tzinfo=cairo_tz).astimezone(ZoneInfo("UTC"))
    pn = Product.name_en if g.current_language == 'en' else Product.name_ar
    ps = ProductSize.size_en if g.current_language == 'en' else ProductSize.size_ar
    q = db.session.query(pn.label("product_name"), ps.label("product_size"),
        db.func.sum(OrderItem.quantity).label("total_quantity"),
        db.func.sum(OrderItem.price * OrderItem.quantity).label("total_price")
    ).select_from(OrderItem).join(Order).join(ProductSize).join(Product)
    if status: q = q.filter(Order.status == status)
    if payment_method: q = q.filter(Order.payment_method == payment_method)
    if start_date and end_date: q = q.filter(Order.date >= start_date, Order.date < end_date_correction(end_date))
    elif start_date: q = q.filter(Order.date >= start_date)
    elif end_date: q = q.filter(Order.date < end_date_correction(end_date))
    q = q.group_by(pn, ps).order_by(pn.asc(), ps.asc())
    pagination = q.paginate(page=page, per_page=50, error_out=False)
    return render_template("admin_reports.html", results=pagination.items, pagination=pagination,
        start_date=start_date, end_date=end_date, status=status, payment_method=payment_method, translations=g.translations)


@admin_reports_bp.route("/admin/reports/download", methods=["GET"], endpoint='download_report')
@login_required
def download_report():
    if current_user.role not in ['admin', 'superadmin']: return redirect(url_for("home"))
    start_date = request.args.get("start_date", None); end_date = request.args.get("end_date", None)
    status = request.args.get("status", ""); payment_method = request.args.get("payment_method", "")
    if start_date: start_date = mydtfmt(start_date).replace(tzinfo=cairo_tz).astimezone(ZoneInfo("UTC"))
    if end_date: end_date = mydtfmt(end_date).replace(tzinfo=cairo_tz).astimezone(ZoneInfo("UTC"))
    pn = Product.name_en if g.current_language == 'en' else Product.name_ar
    ps = ProductSize.size_en if g.current_language == 'en' else ProductSize.size_ar
    q = db.session.query(pn.label("product_name"), ps.label("product_size"),
        db.func.sum(OrderItem.quantity).label("total_quantity"),
        db.func.sum(OrderItem.price * OrderItem.quantity).label("total_price")
    ).select_from(OrderItem).join(Order).join(ProductSize).join(Product)
    if status: q = q.filter(Order.status == status)
    if payment_method: q = q.filter(Order.payment_method == payment_method)
    if start_date and end_date: q = q.filter(Order.date >= start_date, Order.date < end_date_correction(end_date))
    elif start_date: q = q.filter(Order.date >= start_date)
    elif end_date: q = q.filter(Order.date < end_date_correction(end_date))
    results = q.group_by(pn, ps).order_by(pn.asc(), ps.asc()).all()
    df = pd.DataFrame([(r.product_name, r.product_size, r.total_quantity, r.total_price) for r in results],
        columns=[g.translations["Product_Name"], g.translations["Product_Size"], g.translations["Quantity"], f"{g.translations['Total_Price']} ({g.translations['EGP']})"])
    df[g.translations["Payment_Method"]] = g.translations.get(payment_method.replace(' ','_'), g.translations["All"]) if payment_method else g.translations["All"]
    df[g.translations["Status"]] = g.translations.get(status, g.translations["All"]) if status else g.translations["All"]
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as w: df.to_excel(w, index=False, sheet_name="Report")
    output.seek(0)
    return send_file(output, download_name=f"{today()}-report.xlsx", as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@admin_reports_bp.route("/admin/transactions", methods=["GET"], endpoint='admin_transactions')
@login_required
def admin_transactions():
    if current_user.role not in ['admin', 'superadmin']: return redirect(url_for("home"))
    page = request.args.get('page', 1, type=int)
    start_date = request.args.get("start_date", None); end_date = request.args.get("end_date", None)
    success = request.args.get("success", "")
    if start_date: start_date = mydtfmt(start_date).replace(tzinfo=cairo_tz).astimezone(ZoneInfo("UTC"))
    if end_date: end_date = mydtfmt(end_date).replace(tzinfo=cairo_tz).astimezone(ZoneInfo("UTC"))
    q = db.session.query(Transaction)
    if success: q = q.filter(Transaction.success == (success == 'true'))
    if start_date and end_date: q = q.filter(Transaction.updated_at >= start_date, Transaction.updated_at < end_date_correction(end_date))
    elif start_date: q = q.filter(Transaction.updated_at >= start_date)
    elif end_date: q = q.filter(Transaction.updated_at < end_date_correction(end_date))
    q = q.order_by(Transaction.updated_at.desc())
    pagination = q.paginate(page=page, per_page=50, error_out=False)
    return render_template("admin_transactions.html", results=pagination.items, pagination=pagination,
        start_date=start_date, end_date=end_date, success=success)


@admin_reports_bp.route("/admin/transactions/download", methods=["GET"], endpoint='download_transactions')
@login_required
def download_transactions():
    if current_user.role != "admin": return redirect(url_for("home"))
    start_date = request.args.get("start_date", None); end_date = request.args.get("end_date", None)
    success = request.args.get("success", "")
    if start_date: start_date = mydtfmt(start_date).replace(tzinfo=cairo_tz).astimezone(ZoneInfo("UTC"))
    if end_date: end_date = mydtfmt(end_date).replace(tzinfo=cairo_tz).astimezone(ZoneInfo("UTC"))
    q = db.session.query(Transaction)
    if success: q = q.filter(Transaction.success == (success == 'true'))
    if start_date and end_date: q = q.filter(Transaction.updated_at >= start_date, Transaction.updated_at < end_date_correction(end_date))
    elif start_date: q = q.filter(Transaction.updated_at >= start_date)
    elif end_date: q = q.filter(Transaction.updated_at < end_date_correction(end_date))
    results = q.order_by(Transaction.updated_at.desc()).all()
    t = g.translations
    df = pd.DataFrame([(r.id, r.order_id, r.transaction_id, r.amount_cents/100, t[r.currency], r.txn_response_code,
        t["Yes"] if r.success else t["No"], datetimeformat(to_cairo_time(r.created_at)), datetimeformat(to_cairo_time(r.updated_at)),
        t["Yes"] if r.is_3d_secure else t["No"], r.payment_type.upper(), r.payment_sub_type,
        t["Yes"] if r.is_voided else t["No"], t["Yes"] if r.is_refunded else t["No"], r.refunded_amount_cents/100,
        t["Yes"] if r.is_settled else t["No"], r.captured_amount_cents/100, t["Yes"] if r.error_occured else t["No"],
        r.amount_cents/100 - r.refunded_amount_cents/100) for r in results],
        columns=[t['Internal_Trn_ID'], t['Order_ID'], t['PayMob_Trn_ID'], t['Amount'], t['Currency'], t['Trn_Response'],
        t['Success'], t['Created_At'], t['Updated_At'], t['3D_Secured'], t['Payment_Type'], t['Card_Type'],
        t['Voided'], t['Refunded'], t['Refunded_Amount'], t['Settled'], t['Captured_Amount'], t['Error'], t['Net_Amount']])
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as w: df.to_excel(w, index=False, sheet_name="Transactions Report")
    output.seek(0)
    return send_file(output, download_name=f"{today()}-transactions-report.xlsx", as_attachment=True)


@admin_reports_bp.route("/admin/generate_promocodes", methods=["POST"], endpoint='generate_promocodes')
@login_required
def generate_promocodes():
    if current_user.role != "admin":
        flash(g.translations["unauthorized_access"], "danger")
        return redirect(url_for("home"))
    num_codes = int(request.form.get("num_codes", 0))
    discount_percent = float(request.form.get("discount_percent", 0))
    if num_codes <= 0 or discount_percent <= 0 or discount_percent > 100:
        flash(g.translations["invalid_input_please_provide_valid_values"], "danger")
        return redirect(url_for("home"))
    promocodes = []
    for _ in range(num_codes):
        code = ''.join(random.choices(string.ascii_uppercase + string.digits, k=8))
        promocode = PromoCode(code=code, discount_percent=discount_percent)
        promocodes.append(promocode)
    db.session.bulk_save_objects(promocodes)
    db.session.commit()
    flash(g.translations["Promocodes_generated_with_discount"].format(num_codes=num_codes, discount_percent=discount_percent), "success")
    return redirect(url_for("home"))


@admin_reports_bp.route("/admin/download_promocodes", methods=["GET"], endpoint='download_promocodes')
@login_required
def download_promocodes():
    if current_user.role != "admin":
        flash(g.translations["unauthorized_access"], "danger")
        return redirect(url_for("home"))
    promocodes = PromoCode.query.filter_by(is_used=False).all()
    if not promocodes:
        flash(g.translations["no_promocodes_available"], "danger")
        return redirect(url_for("home"))
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(["Promo Code", "Discount Percent"])
    for code in promocodes:
        writer.writerow([code.code, f"{code.discount_percent}%"])
    output.seek(0)
    return send_file(
        BytesIO(output.getvalue().encode()),
        mimetype="text/csv",
        as_attachment=True,
        download_name="promocodes.csv"
    )


@admin_reports_bp.route('/admin/download_products_json', methods=["GET"], endpoint='download_json')
@login_required
def download_json():
    if current_user.role != "admin":
        return redirect(url_for("home"))
    products = Product.query.all()
    products_data = []
    for product in products:
        ingredients_en = product.ingredients_en
        ingredients_ar = product.ingredients_ar
        gender_en = product.gender_en
        gender_ar = product.gender_ar
        products_data.append({
            "name": {
                "en": product.name_en,
                "ar": product.name_ar
            },
            "category": {
                "en": product.category_en,
                "ar": product.category_ar
            },
            "ingredients": {
                "en": ingredients_en,
                "ar": ingredients_ar
            },
            "gender": {
                "en": gender_en,
                "ar": gender_ar
            },
            "description": {
                "en": product.description_en,
                "ar": product.description_ar
            },
            "image_url": product.image_url,
            "image_url_2": product.image_url_2 if product.image_url_2 else None,
            "image_url_3": product.image_url_3 if product.image_url_3 else None,
            "image_url_4": product.image_url_4 if product.image_url_4 else None,
            "image_url_5": product.image_url_5 if product.image_url_5 else None,
            "image_url_6": product.image_url_6 if product.image_url_6 else None,
            "image_url_7": product.image_url_7 if product.image_url_7 else None,
            "image_url_8": product.image_url_8 if product.image_url_8 else None,
            "image_url_9": product.image_url_9 if product.image_url_9 else None,
            "image_url_10": product.image_url_10 if product.image_url_10 else None
        })
    json_data = json.dumps(products_data, ensure_ascii=False, indent=4)
    output = BytesIO()
    output.write(json_data.encode('utf-8'))
    output.seek(0)
    today_str = datetime.now().strftime('%Y-%m-%d')
    return send_file(
        output,
        download_name=f"{today_str}-products_data.json",
        as_attachment=True,
        mimetype='application/json'
    )


@admin_reports_bp.route('/admin/download_products_excel', methods=["GET"], endpoint='download_excel')
@login_required
def download_excel():
    if current_user.role != "admin":
        return redirect(url_for("home"))
    products = Product.query.all()
    products_data = []
    for product in products:
        products_data.append({
            "name_en": product.name_en,
            "name_ar": product.name_ar,
            "category_en": product.category_en,
            "category_ar": product.category_ar,
            "ingredients_en": product.ingredients_en,
            "ingredients_ar": product.ingredients_ar,
            "gender_en": product.gender_en,
            "gender_ar": product.gender_ar,
            "description_en": product.description_en,
            "description_ar": product.description_ar,
            "image_url": product.image_url,
            "image_url_2": product.image_url_2 if product.image_url_2 else None,
            "image_url_3": product.image_url_3 if product.image_url_3 else None,
            "image_url_4": product.image_url_4 if product.image_url_4 else None,
            "image_url_5": product.image_url_5 if product.image_url_5 else None,
            "image_url_6": product.image_url_6 if product.image_url_6 else None,
            "image_url_7": product.image_url_7 if product.image_url_7 else None,
            "image_url_8": product.image_url_8 if product.image_url_8 else None,
            "image_url_9": product.image_url_9 if product.image_url_9 else None,
            "image_url_10": product.image_url_10 if product.image_url_10 else None
        })
    df = pd.DataFrame(products_data)
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Products Data")
    output.seek(0)
    today_str = datetime.now().strftime('%Y-%m-%d')
    return send_file(output, download_name=f"{today_str}-products_data.xlsx", as_attachment=True)


@admin_reports_bp.route('/admin/download_db_files', methods=["GET"], endpoint='download_db_files')
@login_required
def download_db_files():
    if current_user.role != "admin":
        return redirect(url_for("home"))

    db_host = os.getenv('DB_HOST')
    db_name = os.getenv('DB_NAME')
    db_user = os.getenv('DB_USER')
    db_password = os.getenv('DB_PASSWORD')
    output_file = os.getenv('OUTPUT_FILE')
    connection = None
    cursor = None

    try:
        connection = psycopg2.connect(
            host=db_host,
            database=db_name,
            user=db_user,
            password=db_password
        )
        connection.set_client_encoding('UTF8')
        cursor = connection.cursor()
        cursor.execute(
            "SELECT table_name FROM information_schema.tables WHERE table_schema = 'public';"
        )
        tables = cursor.fetchall()
        workbook = Workbook()
        workbook.remove(workbook.active)
        for table_name_tuple in tables:
            table_name = table_name_tuple[0]
            cursor.execute(f"SELECT * FROM \"{table_name}\"")
            rows = cursor.fetchall()
            columns = [desc[0] for desc in cursor.description]
            sheet = workbook.create_sheet(title=table_name)
            for col_num, column_title in enumerate(columns, 1):
                sheet.cell(row=1, column=col_num, value=column_title)

            for row_num, row_data in enumerate(rows, 2):
                for col_num, cell_value in enumerate(row_data, 1):
                    if isinstance(cell_value, list):
                        cell_value = ', '.join(map(str, cell_value))
                    elif isinstance(cell_value, dict):
                        cell_value = str(cell_value)
                    elif isinstance(cell_value, (datetime, dttime)) and cell_value.tzinfo is not None:
                        cell_value = cell_value.replace(tzinfo=None)
                    sheet.cell(row=row_num, column=col_num, value=cell_value)

        workbook.save(output_file)
    except Exception as e:
        print(f"An error occurred: {e}")
    finally:
        if connection:
            if cursor:
                cursor.close()
            connection.close()
    return send_file(output_file, as_attachment=True)

