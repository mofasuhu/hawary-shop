import os
import psycopg2
import openpyxl
from openpyxl import Workbook
from dotenv import load_dotenv
from datetime import datetime, date, time

load_dotenv()

# --- Configuration from .env ---
DB_HOST     = os.getenv('DB_HOST')
DB_NAME     = os.getenv('DB_NAME')
DB_USER     = os.getenv('DB_USER')
DB_PASSWORD = os.getenv('DB_PASSWORD')
DB_PORT     = os.getenv('DB_PORT', '5432')


def strip_timezone(value):
    """Remove timezone info from datetime/time objects so Excel can handle them."""
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    if isinstance(value, time) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def export_db_to_excel(output_file='db_export.xlsx'):
    if not all([DB_HOST, DB_NAME, DB_USER, DB_PASSWORD]):
        print("Error: Database credentials are not fully configured in .env file.")
        return

    print(f"Connecting to database '{DB_NAME}' on '{DB_HOST}'...")

    connection = None
    try:
        connection = psycopg2.connect(
            host=DB_HOST,
            port=DB_PORT,
            database=DB_NAME,
            user=DB_USER,
            password=DB_PASSWORD
        )
        connection.set_client_encoding('UTF8')
        cursor = connection.cursor()

        # Fetch all table names in the public schema — exclude pg_ extension tables
        cursor.execute(
            "SELECT table_name FROM information_schema.tables "
            "WHERE table_schema = 'public' "
            "AND table_name NOT LIKE 'pg_%' "
            "ORDER BY table_name;"
        )
        tables = [row[0] for row in cursor.fetchall()]
        print(f"Found {len(tables)} tables: {', '.join(tables)}")

        workbook = Workbook()
        workbook.remove(workbook.active)  # remove default empty sheet

        for table_name in tables:
            try:
                cursor.execute(f'SELECT * FROM "{table_name}"')
            except Exception as table_err:
                print(f"  ⚠ Skipping '{table_name}': {table_err}")
                connection.rollback()  # reset transaction after error
                continue

            rows = cursor.fetchall()
            columns = [desc[0] for desc in cursor.description]

            sheet = workbook.create_sheet(title=table_name[:31])  # Excel sheet name limit: 31 chars

            # Header row
            for col_num, column_title in enumerate(columns, 1):
                sheet.cell(row=1, column=col_num, value=column_title)

            # Data rows
            for row_num, row_data in enumerate(rows, 2):
                for col_num, cell_value in enumerate(row_data, 1):
                    # Convert lists/dicts to string
                    if isinstance(cell_value, (list, dict)):
                        cell_value = str(cell_value)
                    # Strip timezone info — Excel does not support tz-aware datetimes
                    cell_value = strip_timezone(cell_value)
                    sheet.cell(row=row_num, column=col_num, value=cell_value)

            print(f"  ✔ {table_name}: {len(rows)} rows")

        workbook.save(output_file)
        print(f"\n✅ Exported successfully → {output_file}")

    except Exception as e:
        print(f"\nAn error occurred: {e}")

    finally:
        if connection:
            cursor.close()
            connection.close()


if __name__ == '__main__':
    export_db_to_excel()
